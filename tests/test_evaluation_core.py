"""``plugins/evaluation`` 核心领域层单元测试。

覆盖：models（最大余数/保底分配）、validators（FR6.9）、state（六步校验 +
防抖落盘 + 原子写）、clarification（分级/同类复用/skip/defer/留痕）、
exporters（Excel 合并区 + R5 起始 + 无示例行；Word 段落结构与页眉页脚域；
zip 命名/排序/重名去重；文件名清洗）、template_manager（内置/覆盖/版本/重置）、
importers（知识点/样例题库/考纲）。

不依赖 Flet 运行时；Word 相关用例在缺少 python-docx 时自动跳过。
"""

import io
import json
import shutil
import sys
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PLUGINS_DIR = ROOT / "plugins"
if str(PLUGINS_DIR) not in sys.path:
    sys.path.insert(0, str(PLUGINS_DIR))
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

ROOT = Path(__file__).resolve().parents[1]
PLUGIN_DIR = ROOT / "plugins" / "evaluation"
TEMPLATES_DIR = PLUGIN_DIR / "templates"

try:  # python-docx 是插件依赖，缺失时仅跳过 Word 用例
    import docx  # noqa: F401

    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False

try:
    from openpyxl import load_workbook

    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - openpyxl 是根依赖
    HAS_OPENPYXL = False

from evaluation import importers
from evaluation.clarification import (
    AI_DECIDES,
    STATUS_ANSWERED,
    STATUS_AUTO,
    STATUS_PENDING,
    STATUS_SKIPPED,
    Clarification,
    ClarificationQueue,
    summarize_records,
)
from evaluation.exporter import (
    EXCEL_FIRST_DATA_ROW,
    EXCEL_HEADER_ROW,
    EvaluationExporter,
    format_question_stem,
    knowledge_point_file_stem,
    safe_filename_segment,
    write_exports,
    zip_name,
)
from evaluation.ai_client import (
    AIConfigError,
    StructuredAIClient,
)
from evaluation.models import (
    DEFAULT_DIFFICULTY_RATIO,
    DEFAULT_TYPE_RATIO,
    Chapter,
    CourseDraft,
    KnowledgePoint,
    Question,
    allocate_by_largest_remainder,
    allocate_with_floor,
    ratio_from_counts,
)
from evaluation.state import (
    DraftStore,
    WizardValidator,
    apply_global_defaults,
    apply_global_details,
    pending_clarifications,
)
from evaluation.prompts import build_generation_request
from evaluation.template_manager import TEMPLATE_FILES, TemplateManager
from evaluation.validators import sort_questions, validate_question, validate_questions


def make_point(**kwargs) -> KnowledgePoint:
    defaults = dict(
        name="窗口结构",
        order=1,
        chapter_order=1,
        chapter_name="界面介绍",
        level="对口",
        knowledge_type="了解",
        tags=["标题栏"],
    )
    defaults.update(kwargs)
    return KnowledgePoint(**defaults)


def make_questions() -> list:
    return [
        Question(
            stem="CorelDraw窗口界面的组成（）",
            question_type="单选",
            difficulty="一般",
            options=["标题栏、菜单栏", "属性栏", "网格", "色盘"],
            answer="A",
            confidence=0.9,
            internal_id="1-1-X-1",
        ),
        Question(
            stem="CorelDRAW可链接的软件有（）",
            question_type="多选",
            difficulty="困难",
            options=["Dreamweaver", "Photoshop", "Illustrator", "Flash"],
            answer="BC",
            confidence=0.7,
            internal_id="1-1-D-1",
        ),
        Question(
            stem="图片插入后可通过拖动边角调整大小。",
            question_type="判断",
            difficulty="简单",
            options=["正确", "错误"],
            answer="A",
            confidence=0.5,
            internal_id="1-1-P-1",
        ),
    ]


# ==================================================================== models
class LargestRemainderTests(unittest.TestCase):
    def test_sum_always_equals_total(self):
        for total in list(range(0, 60)) + [100, 101]:
            with self.subTest(total=total):
                counts = allocate_by_largest_remainder(
                    total, {"单选": 70, "多选": 20, "判断": 10}
                )
                self.assertEqual(sum(counts.values()), total)

    def test_default_ratio_split(self):
        self.assertEqual(
            allocate_by_largest_remainder(10, {"单选": 70, "多选": 20, "判断": 10}),
            {"单选": 7, "多选": 2, "判断": 1},
        )

    def test_zero_total_and_zero_weights(self):
        self.assertEqual(
            allocate_by_largest_remainder(0, {"单选": 70}), {"单选": 0}
        )
        self.assertEqual(
            allocate_by_largest_remainder(5, {"单选": 0, "多选": 0}),
            {"单选": 0, "多选": 0},
        )

    def test_floor_gives_every_selected_type_at_least_one(self):
        for total in range(3, 40):
            with self.subTest(total=total):
                counts = allocate_with_floor(
                    total, {"单选": 70, "多选": 20, "判断": 10}
                )
                self.assertEqual(sum(counts.values()), total)
                self.assertTrue(all(value >= 1 for value in counts.values()))

    def test_floor_falls_back_when_total_too_small(self):
        counts = allocate_with_floor(2, {"单选": 70, "多选": 20, "判断": 10})
        self.assertEqual(sum(counts.values()), 2)

    def test_ratio_from_counts_sums_to_100(self):
        ratio = ratio_from_counts({"单选": 3, "多选": 2, "判断": 1})
        self.assertEqual(sum(ratio.values()), 100)
        self.assertEqual(ratio_from_counts({"单选": 0, "多选": 0, "判断": 0}), {"单选": 0, "多选": 0, "判断": 0})


class KnowledgePointDistributionTests(unittest.TestCase):
    def test_fresh_point_is_valid_by_default(self):
        point = KnowledgePoint(name="K")
        self.assertEqual(sum(point.type_counts.values()), point.question_count)
        self.assertEqual(point.selected_ratio_total, 100)
        self.assertEqual(
            WizardValidator().validate_step(
                CourseDraft(
                    course_name="C",
                    chapters=[Chapter(name="ch", knowledge_points=[point])],
                ),
                4,
            ),
            [],
        )

    def test_deselecting_a_type_zeroes_it_and_normalizes_ratio(self):
        point = KnowledgePoint(name="K")
        point.set_type_ratio({"单选": 80, "多选": 20, "判断": 0})
        self.assertEqual(point.type_ratio["判断"], 0)
        self.assertEqual(point.selected_ratio_total, 100)
        self.assertEqual(point.type_counts["判断"], 0)

    def test_counts_to_ratio_roundtrip(self):
        point = KnowledgePoint(name="K")
        point.set_type_counts({"单选": 5, "多选": 2, "判断": 0})
        self.assertEqual(point.selected_ratio_total, 100)
        self.assertEqual(point.type_counts, {"单选": 5, "多选": 2, "判断": 0})

    def test_set_question_count_keeps_sum(self):
        point = KnowledgePoint(name="K")
        for total in (3, 10, 30):
            point.set_question_count(total)
            self.assertEqual(sum(point.type_counts.values()), total)

    def test_legacy_ratio_in_type_counts_is_migrated(self):
        point = KnowledgePoint(
            name="K", question_count=10, type_counts=dict(DEFAULT_TYPE_RATIO)
        )
        self.assertEqual(sum(point.type_counts.values()), 10)

    def test_defaults_constants_add_up(self):
        self.assertEqual(sum(DEFAULT_TYPE_RATIO.values()), 100)
        self.assertEqual(sum(DEFAULT_DIFFICULTY_RATIO.values()), 100)


class CourseDraftTests(unittest.TestCase):
    def _draft(self) -> CourseDraft:
        point = make_point()
        point.questions = make_questions()
        draft = CourseDraft(
            course_name="CorelDraw",
            chapters=[Chapter(name="界面介绍", knowledge_points=[point])],
        )
        draft.renumber()
        return draft

    def test_roundtrip_is_identical(self):
        draft = self._draft()
        restored = CourseDraft.from_dict(draft.to_dict())
        self.assertEqual(restored.to_dict(), draft.to_dict())

    def test_renumber_fills_chapter_and_point_order(self):
        draft = CourseDraft(
            course_name="C",
            chapters=[
                Chapter(name="A", knowledge_points=[KnowledgePoint(name="a1")]),
                Chapter(name="B", knowledge_points=[KnowledgePoint(name="b1"), KnowledgePoint(name="b2")]),
            ],
        )
        draft.renumber()
        self.assertEqual([p.order for p in draft.all_points], [1, 1, 2])
        self.assertEqual([p.chapter_order for p in draft.all_points], [1, 2, 2])
        self.assertEqual(draft.all_points[2].chapter_name, "B")

    def test_total_questions(self):
        self.assertEqual(self._draft().total_questions, 10)

    def test_confidence_colour_bands(self):
        self.assertEqual(Question(confidence=0.8).confidence_color, "green")
        self.assertEqual(Question(confidence=0.79).confidence_color, "amber")
        self.assertEqual(Question(confidence=0.6).confidence_color, "amber")
        self.assertEqual(Question(confidence=0.59).confidence_color, "red")

    def test_effective_tags_prefers_teacher_tags(self):
        point = KnowledgePoint(name="K", tags=["t"], inferred_tags=["i"])
        self.assertEqual(point.effective_tags, ["t"])
        point.tags = []
        self.assertEqual(point.effective_tags, ["i"])


# ================================================================ validators
class ValidatorTests(unittest.TestCase):
    def test_valid_questions_pass(self):
        for question in make_questions():
            with self.subTest(question=question.question_type):
                self.assertEqual(validate_question(question), [])

    def test_blank_stem_rejected(self):
        self.assertIn("题干不能为空", validate_question(Question(stem="  ")))

    def test_bad_type_and_difficulty_rejected(self):
        errors = validate_question(
            Question(stem="x", question_type="填空", difficulty="超难", options=["a"], answer="A")
        )
        self.assertTrue(any("题型枚举不合规" in item for item in errors))
        self.assertTrue(any("难度枚举不合规" in item for item in errors))

    def test_more_than_eight_options_rejected(self):
        errors = validate_question(
            Question(stem="x", options=[str(i) for i in range(9)], answer="A")
        )
        self.assertTrue(any("选项不能超过8个" in item for item in errors))

    def test_multi_choice_requires_two_letters(self):
        errors = validate_question(
            Question(stem="x", question_type="多选", options=["a", "b"], answer="A")
        )
        self.assertIn("多选答案至少包含两个字母", errors)

    def test_single_choice_requires_one_letter(self):
        errors = validate_question(
            Question(stem="x", question_type="单选", options=["a", "b", "c"], answer="AB")
        )
        self.assertIn("单选答案只能有一个字母", errors)

    def test_judge_options_and_answer_are_pinned(self):
        errors = validate_question(
            Question(stem="x", question_type="判断", options=["对", "错"], answer="A")
        )
        self.assertIn("判断题选项必须为A.正确、B.错误", errors)

        errors = validate_question(
            Question(stem="x", question_type="判断", options=["正确", "错误"], answer="C")
        )
        self.assertIn("判断题答案必须为A或B", errors)

    def test_answer_letter_must_map_to_existing_option(self):
        errors = validate_question(
            Question(stem="x", question_type="单选", options=["a", "b"], answer="C")
        )
        self.assertTrue(any("没有对应选项" in item for item in errors))

    def test_duplicate_answer_letters_rejected(self):
        errors = validate_question(
            Question(stem="x", question_type="多选", options=["a", "b"], answer="AA")
        )
        self.assertTrue(any("答案字母不能重复" in item for item in errors))

    def test_validate_questions_locates_problem(self):
        errors = validate_questions(
            [Question(stem="x", question_type="多选", options=["a", "b"], answer="A", internal_id="1-2-D-3")]
        )
        self.assertEqual(len(errors), 1)
        self.assertTrue(errors[0].startswith("1-2-D-3："))

    def test_sort_questions_orders_by_type_then_number(self):
        questions = [
            Question(stem="p2", question_type="判断", internal_id="1-1-P-2"),
            Question(stem="x2", question_type="单选", internal_id="1-1-X-2"),
            Question(stem="d", question_type="多选", internal_id="1-1-D-1"),
            Question(stem="x10", question_type="单选", internal_id="1-1-X-10"),
            Question(stem="x1", question_type="单选", internal_id="1-1-X-1"),
        ]
        ordered = [q.internal_id for q in sort_questions(questions)]
        self.assertEqual(ordered, ["1-1-X-1", "1-1-X-2", "1-1-X-10", "1-1-D-1", "1-1-P-2"])


# ================================================================== state
class DraftStoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.clock = [0.0]
        self.store = DraftStore(
            self.tmp, min_interval=2.0, every_n_writes=5, time_fn=lambda: self.clock[0]
        )
        self.draft = CourseDraft(
            course_name="X",
            chapters=[Chapter(name="C", knowledge_points=[KnowledgePoint(name="K")])],
        )
        self.draft.renumber()

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_load_missing_file_returns_fresh_draft(self):
        self.assertEqual(self.store.load().course_name, "")
        self.assertEqual(self.store.load().chapters, [])

    def test_load_corrupted_file_returns_fresh_draft(self):
        self.store.path.parent.mkdir(parents=True, exist_ok=True)
        self.store.path.write_text("{not json", encoding="utf-8")
        self.assertEqual(self.store.load().course_name, "")

    def test_flush_without_dirty_is_noop(self):
        self.assertFalse(self.store.flush(self.draft))
        self.assertFalse(self.store.flush_if_due(self.draft))
        self.assertEqual(self.store.flush_count, 0)

    def test_debounce_waits_until_threshold(self):
        self.store.mark_dirty()
        self.assertFalse(self.store.flush_if_due(self.draft))
        self.assertEqual(self.store.flush_count, 0)

        for _ in range(4):
            self.store.mark_dirty()
        self.assertTrue(self.store.flush_if_due(self.draft))
        self.assertEqual(self.store.flush_count, 1)

    def test_debounce_flushes_after_time_interval(self):
        self.store.mark_dirty()
        self.assertFalse(self.store.flush_if_due(self.draft))
        self.clock[0] = 3.0
        self.assertTrue(self.store.flush_if_due(self.draft))
        self.assertEqual(self.store.flush_count, 1)

    def test_atomic_write_leaves_no_tmp_file(self):
        self.store.mark_dirty()
        self.assertTrue(self.store.flush(self.draft, force=True))
        self.assertTrue(self.store.path.is_file())
        self.assertFalse(self.store.path.with_suffix(self.store.path.suffix + ".tmp").exists())

    def test_roundtrip_and_clear(self):
        self.store.flush(self.draft, force=True)
        self.assertEqual(self.store.load().course_name, "X")
        self.store.clear()
        self.assertFalse(self.store.path.exists())
        self.assertEqual(self.store.load().course_name, "")


class WizardValidatorTests(unittest.TestCase):
    def setUp(self):
        self.validator = WizardValidator()

    def _draft(self, **kwargs) -> CourseDraft:
        point = make_point()
        draft = CourseDraft(
            course_name="CorelDraw",
            chapters=[Chapter(name="界面介绍", knowledge_points=[point])],
            **kwargs,
        )
        draft.renumber()
        return draft

    def test_step_one_requires_course_name(self):
        self.assertEqual(
            self.validator.validate_step(CourseDraft(), 1), ["课程名称不能为空"]
        )
        self.assertEqual(self.validator.validate_step(self._draft(), 1), [])

    def test_step_two_requires_chapter_and_point(self):
        self.assertIn(
            "至少创建一个章节", self.validator.validate_step(CourseDraft(course_name="C"), 2)
        )
        draft = CourseDraft(course_name="C", chapters=[Chapter(name="A")])
        self.assertIn(
            "章节“A”至少需要一个知识点", self.validator.validate_step(draft, 2)
        )
        self.assertEqual(self.validator.validate_step(self._draft(), 2), [])

    def test_step_three_never_blocks(self):
        self.assertEqual(self.validator.validate_step(CourseDraft(), 3), [])

    def test_step_four_rules(self):
        draft = self._draft()
        point = draft.all_points[0]
        self.assertEqual(self.validator.validate_step(draft, 4), [])

        point.question_count = 2
        errors = self.validator.validate_step(draft, 4)
        self.assertTrue(any("题型题数合计必须等于总题数" in item for item in errors))

        point.question_count = 10
        point.difficulty_ratio = {"简单": 30, "一般": 30, "困难": 30}
        self.assertTrue(
            any("难度比例必须合计100%" in item for item in self.validator.validate_step(draft, 4))
        )

    def test_step_four_flags_too_few_questions_for_selected_types(self):
        draft = self._draft()
        point = draft.all_points[0]
        point.set_type_ratio({"单选": 34, "多选": 33, "判断": 33}, question_count=2)
        errors = self.validator.validate_step(draft, 4)
        self.assertTrue(any("题数不足以覆盖已选题型" in item for item in errors))

    def test_step_five_requires_complete_and_no_pending(self):
        draft = self._draft()
        self.assertIn("生成尚未完成", self.validator.validate_step(draft, 5))

        draft.generation_status = "complete"
        self.assertEqual(self.validator.validate_step(draft, 5), [])

        draft.clarification_records.append({"status": "pending", "question": "q"})
        errors = self.validator.validate_step(draft, 5)
        self.assertTrue(any("待澄清" in item for item in errors))
        self.assertEqual(len(pending_clarifications(draft)), 1)

    def test_step_six_needs_at_least_one_exportable_point(self):
        self.assertIn("至少勾选1个知识点", self.validator.validate_step(CourseDraft(), 6))
        draft = self._draft()
        self.assertEqual(self.validator.validate_step(draft, 6), [])
        draft.all_points[0].skipped = True
        self.assertIn("至少勾选1个知识点", self.validator.validate_step(draft, 6))


class ApplyGlobalTests(unittest.TestCase):
    def _draft(self, count=2) -> CourseDraft:
        draft = CourseDraft(
            course_name="C",
            chapters=[
                Chapter(
                    name="A",
                    knowledge_points=[KnowledgePoint(name=f"P{i}") for i in range(count)],
                )
            ],
        )
        draft.renumber()
        return draft

    def test_apply_global_defaults_sets_all_rows(self):
        draft = self._draft()
        affected = apply_global_defaults(draft, question_count=12)
        self.assertEqual(affected, 2)
        for point in draft.all_points:
            self.assertEqual(point.question_count, 12)
            self.assertEqual(sum(point.type_counts.values()), 12)
            self.assertFalse(point.overridden)
        self.assertEqual(WizardValidator().validate_step(draft, 4), [])

    def test_apply_global_defaults_can_target_subset(self):
        draft = self._draft()
        apply_global_defaults(draft, question_count=7, targets=[draft.all_points[0]])
        self.assertEqual(draft.all_points[0].question_count, 7)
        self.assertEqual(draft.all_points[1].question_count, 10)

    def test_apply_global_details_only_fills_unfilled_tags(self):
        draft = self._draft()
        draft.all_points[1].tags = ["existing"]
        affected = apply_global_details(
            draft, level="明星", knowledge_type="掌握", tags=["t1"]
        )
        self.assertEqual(affected, 2)
        self.assertEqual(draft.all_points[0].tags, ["t1"])
        self.assertEqual(draft.all_points[1].tags, ["existing"])
        self.assertTrue(all(point.level == "明星" for point in draft.all_points))


# ============================================================ clarification
class ClarificationQueueTests(unittest.TestCase):
    def _queue(self, sensitivity="标准", records=None) -> ClarificationQueue:
        sink = records if records is not None else []
        return ClarificationQueue(sensitivity=sensitivity, on_record=sink.append)

    def test_sensitivity_gating(self):
        loose = self._queue("宽松")
        confirm = Clarification(kind="k1", level="确认", question="q")
        blocking = Clarification(kind="k2", level="阻断", question="q")
        loose.add(confirm)
        loose.add(blocking)
        self.assertEqual(confirm.status, STATUS_AUTO)
        self.assertEqual(confirm.source, "auto")
        self.assertEqual(blocking.status, STATUS_PENDING)

        strict = self._queue("严格")
        auto = Clarification(kind="k3", level="自决", question="q")
        strict.add(auto)
        self.assertEqual(auto.status, STATUS_PENDING)

    def test_automatic_decisions_are_recorded(self):
        records = []
        queue = self._queue("宽松", records)
        queue.add(Clarification(kind="k", level="自决", question="润色"))
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["choice"], AI_DECIDES)
        self.assertEqual(records[0]["status"], STATUS_AUTO)
        for key in ("question", "choice", "scope", "timestamp", "level"):
            self.assertIn(key, records[0])

    def test_same_kind_reuse_skips_dialog(self):
        records = []
        queue = self._queue("标准", records)
        first = Clarification(kind="version", level="确认", question="q", point_name="P1")
        queue.add(first)
        queue.answer(first, "不敏感", apply_to_kind=True)

        second = Clarification(kind="version", level="确认", question="q", point_name="P2")
        queue.add(second)
        self.assertEqual(second.status, STATUS_AUTO)
        self.assertEqual(second.source, "reuse")
        self.assertEqual(second.choice, "不敏感")
        self.assertEqual(queue.pending(), [])

    def test_worker_is_woken_by_answer(self):
        queue = self._queue()
        clarification = Clarification(kind="k", level="阻断", question="q")
        queue.add(clarification)

        outcome = {}

        def worker():
            outcome["waited"] = clarification.wait(timeout=5)
            outcome["choice"] = clarification.choice

        thread = threading.Thread(target=worker, daemon=True)
        thread.start()
        time.sleep(0.05)
        queue.answer(clarification, "按推荐处理")
        thread.join(timeout=5)

        self.assertTrue(outcome.get("waited"))
        self.assertEqual(outcome.get("choice"), "按推荐处理")
        self.assertFalse(thread.is_alive())

    def test_defer_keeps_pending_but_does_not_block_worker(self):
        queue = self._queue()
        clarification = Clarification(kind="k", level="阻断", question="q")
        queue.add(clarification)

        released = threading.Event()

        def worker():
            clarification.wait(timeout=5)
            released.set()

        thread = threading.Thread(target=worker, daemon=True)
        thread.start()
        time.sleep(0.05)
        queue.defer(clarification)
        thread.join(timeout=5)

        self.assertTrue(released.is_set(), "defer 必须唤醒 worker，否则死循环")
        self.assertEqual(clarification.status, STATUS_PENDING)
        self.assertEqual(len(queue.pending()), 1)

    def test_skip_marks_point_and_records(self):
        records = []
        queue = self._queue("标准", records)
        clarification = Clarification(kind="k", level="阻断", question="q", point_name="P1")
        queue.add(clarification)
        queue.skip(clarification)
        queue.mark_point_skipped(1, 2)

        self.assertEqual(clarification.status, STATUS_SKIPPED)
        self.assertEqual(queue.pending(), [])
        self.assertTrue(queue.is_point_skipped(1, 2))
        self.assertEqual(records[-1]["status"], STATUS_SKIPPED)

    def test_counter_and_release_all(self):
        queue = self._queue()
        first = Clarification(kind="a", level="阻断", question="q")
        second = Clarification(kind="b", level="阻断", question="q")
        queue.add(first)
        queue.add(second)
        self.assertEqual(queue.counter_text(), "澄清 0/2")
        queue.answer(first, "x")
        self.assertEqual(queue.counter_text(), "澄清 1/2")

        queue.release_all()
        self.assertTrue(second.event_released)
        self.assertEqual(second.status, STATUS_PENDING)

    def test_summarize_records(self):
        records = [
            {"status": STATUS_ANSWERED},
            {"status": STATUS_AUTO},
            {"status": STATUS_SKIPPED},
        ]
        summary = summarize_records(records)
        self.assertEqual(summary["total"], 3)
        self.assertEqual(summary[STATUS_ANSWERED], 1)


# ================================================================ importer
class ImporterTests(unittest.TestCase):
    def test_import_knowledge_points_dedupes_and_keeps_order(self):
        if not HAS_OPENPYXL:  # pragma: no cover
            self.skipTest("openpyxl 不可用")
        from openpyxl import Workbook

        tmp = Path(tempfile.mkdtemp())
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["知识点名称"])
            sheet.append(["窗口界面"])
            sheet.append(["工具栏"])
            sheet.append(["工具栏"])
            sheet.append([None])
            sheet.append(["文件操作"])
            path = tmp / "kp.xlsx"
            workbook.save(path)

            self.assertEqual(
                importers.import_knowledge_points_xlsx(path),
                ["窗口界面", "工具栏", "文件操作"],
            )
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_import_excel_samples_skips_note_row(self):
        path = TEMPLATES_DIR / TEMPLATE_FILES["excel"]
        samples = importers.import_question_examples(path)
        self.assertEqual(len(samples), 3)
        for sample in samples:
            self.assertNotIn("请删除", sample["stem"])
            self.assertIn(sample["type"], ("单选", "多选", "判断"))
        self.assertEqual([s["type"] for s in samples], ["单选", "多选", "判断"])
        self.assertEqual(samples[1]["answer"], "ABC")
        self.assertEqual(samples[2]["options"], ["正确", "错误"])

    def test_syllabus_text_accepts_xlsx_and_docx(self):
        text = importers.extract_syllabus_text(TEMPLATES_DIR / TEMPLATE_FILES["excel"])
        self.assertIn("注意事项", text)

    def test_syllabus_text_rejects_pdf(self):
        with self.assertRaises(ValueError):
            importers.extract_syllabus_text("some.pdf")

    def test_parse_question_text_block(self):
        block = "\n".join(
            [
                "CorelDraw窗口界面的组成（）【一般】【单选】",
                "A.标题栏、菜单栏",
                "B.属性栏",
                "答案：【B】",
            ]
        )
        parsed = importers.parse_question_text_block(block)
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["answer"], "B")
        self.assertEqual(parsed[0]["type"], "单选")
        self.assertEqual(parsed[0]["options"], ["标题栏、菜单栏", "属性栏"])

    def test_candidate_knowledge_parsing(self):
        self.assertEqual(
            importers.parse_candidate_knowledge_response(
                {"knowledge_points": ["A", "B", "A"]}
            ),
            ["A", "B"],
        )
        self.assertEqual(importers.parse_candidate_knowledge_response("nope"), [])
        self.assertIn(
            "knowledge_points",
            importers.build_candidate_knowledge_prompt("考纲文本", "课程"),
        )

    def test_json_sample_bank(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            path = tmp / "bank.json"
            path.write_text(
                json.dumps(
                    {
                        "课程": {
                            "章节": {
                                "知识点": [
                                    {"question": "q1", "answer": "A", "A": "a1", "B": "b1"},
                                    {"question": "q2", "answer": "BC"},
                                ]
                            }
                        }
                    }
                ),
                encoding="utf-8",
            )
            samples = importers.import_question_examples(path)
            self.assertEqual(len(samples), 2)
            self.assertEqual(samples[0]["options"], ["a1", "b1"])
            self.assertEqual(samples[1]["type"], "多选")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ========================================================= template_manager
class TemplateManagerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        shutil.copytree(PLUGIN_DIR, self.tmp / "evaluation", dirs_exist_ok=True)
        self.manager = TemplateManager(self.tmp / "evaluation")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_bundled_templates_resolve(self):
        for kind in TEMPLATE_FILES:
            resource = self.manager.get(kind)
            self.assertEqual(resource.source, "bundled")
            self.assertTrue(resource.path.is_file())

    def test_install_override_takes_precedence_and_records_version(self):
        resource = self.manager.install_override(
            "excel", TEMPLATES_DIR / TEMPLATE_FILES["excel"], version="2026.1"
        )
        self.assertTrue(resource.is_override)
        self.assertEqual(resource.version, "2026.1")

        current = self.manager.get("excel")
        self.assertEqual(current.source, "override")
        self.assertEqual(current.version, "2026.1")
        self.assertEqual(self.manager.get("word").source, "bundled")

    def test_auto_version_is_generated(self):
        resource = self.manager.install_override(
            "word", TEMPLATES_DIR / TEMPLATE_FILES["word"]
        )
        self.assertTrue(resource.version.isdigit())

    def test_suffix_and_kind_guards(self):
        with self.assertRaises(ValueError):
            self.manager.install_override("excel", TEMPLATES_DIR / TEMPLATE_FILES["word"])
        with self.assertRaises(ValueError):
            self.manager.install_override("pdf", TEMPLATES_DIR / TEMPLATE_FILES["excel"])
        with self.assertRaises(FileNotFoundError):
            self.manager.install_override("excel", self.tmp / "missing.xlsx")

    def test_reset_override_restores_bundled(self):
        self.manager.install_override("excel", TEMPLATES_DIR / TEMPLATE_FILES["excel"])
        self.manager.reset_override("excel")
        self.assertEqual(self.manager.get("excel").source, "bundled")
        self.assertFalse(self.manager.metadata_path.exists())

    def test_missing_bundled_template_raises(self):
        empty = TemplateManager(self.tmp / "empty")
        with self.assertRaises(FileNotFoundError):
            empty.get("word")
        self.assertEqual(empty.available(), {"word": None, "excel": None})


# ================================================================ exporter
class FileNameTests(unittest.TestCase):
    def test_illegal_characters_are_sanitised(self):
        self.assertEqual(
            safe_filename_segment('A/B:C*D?E"F<G>H|I'), "A_B_C_D_E_F_G_H_I"
        )

    def test_reserved_names_are_prefixed(self):
        self.assertEqual(safe_filename_segment("CON"), "_CON")
        self.assertEqual(safe_filename_segment("lpt1"), "_lpt1")

    def test_long_names_are_truncated(self):
        self.assertLessEqual(len(safe_filename_segment("x" * 200)), 80)

    def test_blank_falls_back(self):
        self.assertEqual(safe_filename_segment("   "), "evaluation")
        self.assertEqual(safe_filename_segment("..."), "evaluation")

    def test_stem_and_zip_names(self):
        point = make_point()
        self.assertEqual(knowledge_point_file_stem("CorelDraw", point), "CorelDraw_1_界面介绍_1_窗口结构")
        self.assertEqual(zip_name("CorelDraw", "20260920"), "CorelDraw_评估题库_20260920.zip")

    def test_zip_name_sanitises_illegal_course_name(self):
        self.assertEqual(zip_name("A/B:C", "20260920"), "A_B_C_评估题库_20260920.zip")

    def test_stem_format_matches_template(self):
        self.assertEqual(
            format_question_stem(Question(stem="组成", difficulty="一般", question_type="单选")),
            "组成（）【一般】【单选】",
        )
        self.assertEqual(
            format_question_stem(
                Question(stem="尺寸是（ ）。", difficulty="简单", question_type="多选")
            ),
            "尺寸是（ ）。【简单】【多选】",
        )

    def test_judge_stem_has_no_answer_slot_brackets(self):
        """判断题是陈述句，模板样例不带 （）；模型多写的尾部 （） 要去掉。"""
        # 无括号 → 不补
        self.assertEqual(
            format_question_stem(
                Question(
                    stem="图片插入后可通过拖动边角调整大小。",
                    difficulty="简单",
                    question_type="判断",
                )
            ),
            "图片插入后可通过拖动边角调整大小。【简单】【判断】",
        )
        # 模型多写了全角/半角尾部括号 → 去掉
        self.assertEqual(
            format_question_stem(
                Question(
                    stem="标准工具栏位于菜单栏下方。（）",
                    difficulty="一般",
                    question_type="判断",
                )
            ),
            "标准工具栏位于菜单栏下方。【一般】【判断】",
        )
        self.assertEqual(
            format_question_stem(
                Question(stem="该说法正确。( )", difficulty="一般", question_type="判断")
            ),
            "该说法正确。【一般】【判断】",
        )
        # 中间出现括号时不能误删
        self.assertEqual(
            format_question_stem(
                Question(
                    stem="按F9（放大）可放大视图。",
                    difficulty="一般",
                    question_type="判断",
                )
            ),
            "按F9（放大）可放大视图。【一般】【判断】",
        )


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl 不可用")
class ExcelExporterTests(unittest.TestCase):
    def setUp(self):
        self.exporter = EvaluationExporter(TEMPLATE_DIR_MANAGER)
        self.point = make_point()
        self.point.questions = make_questions()

    def _load(self):
        return load_workbook(io.BytesIO(self.exporter.export_excel(self.point)))

    def test_workbook_keeps_all_sheets(self):
        self.assertEqual(self._load().sheetnames, ["Sheet1", "Sheet2", "Sheet3"])

    def test_meta_merges_and_values(self):
        sheet = self._load()["Sheet1"]
        merges = {str(item) for item in sheet.merged_cells.ranges}
        for expected in ("A1:L1", "A2:D2", "E2:H2", "I2:L2", "A3:D3", "E3:H3", "I3:L3"):
            self.assertIn(expected, merges)

        self.assertTrue(str(sheet["A1"].value).startswith("注意事项"))
        self.assertEqual(sheet["A3"].value, "窗口结构")
        self.assertEqual(sheet["E3"].value, "对口")
        self.assertEqual(sheet["I3"].value, "了解")

    def test_header_row_intact(self):
        sheet = self._load()["Sheet1"]
        headers = [sheet.cell(row=EXCEL_HEADER_ROW, column=col).value for col in range(1, 14)]
        self.assertEqual(
            headers,
            [
                "序号",
                "题目",
                "难度(简单、一般、困难)",
                "类型（单选、多选、判断）",
                "正确答案项",
                "答案A",
                "答案B",
                "答案C",
                "答案D",
                "答案E",
                "答案F",
                "答案G",
                "答案H",
            ],
        )

    def test_data_rows_start_at_r5_without_template_samples(self):
        sheet = self._load()["Sheet1"]
        self.assertEqual(sheet.max_row, EXCEL_FIRST_DATA_ROW + 2)

        first = [sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=col).value for col in range(1, 8)]
        self.assertEqual(
            first,
            [1, "CorelDraw窗口界面的组成（）", "一般", "单选", "A", "标题栏、菜单栏", "属性栏"],
        )
        self.assertEqual(sheet.cell(row=EXCEL_FIRST_DATA_ROW + 1, column=5).value, "BC")
        self.assertEqual(sheet.cell(row=EXCEL_FIRST_DATA_ROW + 2, column=6).value, "正确")

    def test_note_row_is_gone(self):
        sheet = self._load()["Sheet1"]
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                self.assertNotIn("请删除", str(value or ""))

    def test_template_font_is_preserved(self):
        sheet = self._load()["Sheet1"]
        cell = sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=2)
        self.assertEqual(cell.font.name, "宋体")

    def test_repaired_when_template_merges_are_mangled(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            shutil.copytree(PLUGIN_DIR, tmp / "evaluation", dirs_exist_ok=True)
            target = tmp / "evaluation" / "templates" / TEMPLATE_FILES["excel"]
            workbook = load_workbook(target)
            sheet = workbook["Sheet1"]
            for coord in ("A3:D3", "E3:H3", "I3:L3"):
                sheet.unmerge_cells(coord)
            sheet.merge_cells("A3:B3")
            workbook.save(target)

            exporter = EvaluationExporter(TemplateManager(tmp / "evaluation"))
            out = load_workbook(io.BytesIO(exporter.export_excel(self.point)))["Sheet1"]
            merges = {str(item) for item in out.merged_cells.ranges}
            self.assertIn("A3:D3", merges)
            self.assertNotIn("A3:B3", merges)
            self.assertEqual(out["A3"].value, "窗口结构")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_word_and_excel_judge_stems_agree(self):
        """回归：真实跑出的产物里 Excel 判断题还带 （）而 Word 已去掉。

        两份产物必须一致：判断题都不带作答位括号。
        """
        point = make_point()
        point.questions = [
            Question(
                stem="标尺默认显示在顶部和左侧（）",
                question_type="判断",
                difficulty="一般",
                options=["正确", "错误"],
                answer="A",
                internal_id="1-1-P-1",
            )
        ]

        sheet = load_workbook(
            io.BytesIO(self.exporter.export_excel(point))
        )["Sheet1"]
        excel_stem = sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=2).value
        self.assertEqual(excel_stem, "标尺默认显示在顶部和左侧")
        self.assertNotIn("（）", excel_stem)

        if HAS_DOCX:
            document = docx.Document(io.BytesIO(self.exporter.export_word(point)))
            word_stem = [
                p.text for p in document.paragraphs if p.text.endswith("【判断】")
            ][0]
            self.assertEqual(word_stem, "标尺默认显示在顶部和左侧【一般】【判断】")
            self.assertNotIn("（）", word_stem)

    def test_excel_keeps_choice_stem_brackets(self):
        """单选题题干自带的作答位括号在 Excel 里必须保留（那是题干的一部分）。"""
        point = make_point()
        point.questions = [
            Question(
                stem="属性栏的作用是（）",
                question_type="单选",
                difficulty="一般",
                options=["a", "b", "c", "d"],
                answer="A",
                internal_id="1-1-X-1",
            )
        ]
        sheet = load_workbook(
            io.BytesIO(self.exporter.export_excel(point))
        )["Sheet1"]
        self.assertEqual(
            sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=2).value, "属性栏的作用是（）"
        )

    def test_invalid_questions_block_export(self):
        bad = make_point()
        bad.questions = [
            Question(
                stem="x",
                question_type="多选",
                options=["a", "b"],
                answer="A",
                internal_id="1-1-D-1",
            )
        ]
        with self.assertRaises(ValueError) as ctx:
            self.exporter.export_excel(bad)
        self.assertIn("1-1-D-1", str(ctx.exception))


@unittest.skipUnless(HAS_DOCX, "python-docx 不可用")
class WordExporterTests(unittest.TestCase):
    def setUp(self):
        self.exporter = EvaluationExporter(TEMPLATE_DIR_MANAGER)
        self.point = make_point()
        self.point.questions = make_questions()
        self.data = self.exporter.export_word(self.point)

    def _document(self):
        return docx.Document(io.BytesIO(self.data))

    def _xml(self, name: str) -> str:
        return zipfile.ZipFile(io.BytesIO(self.data)).read(name).decode("utf-8")

    def test_header_title_is_replaced(self):
        header = self._document().sections[0].header
        runs = [run.text for run in header.paragraphs[0].runs]
        self.assertIn("窗口结构", runs)
        self.assertIn("”知识点评估题", runs)
        self.assertNotIn("CorelDraw界面介绍", runs)

    def test_footer_keeps_page_fields(self):
        footer = self._document().sections[0].footer
        text = "".join(run.text for run in footer.paragraphs[0].runs)
        self.assertIn("第", text)
        self.assertIn("页 共", text)

        xml = self._xml("word/footer1.xml")
        self.assertIn("PAGE", xml)
        self.assertIn("NUMPAGES", xml)

    def test_body_layout(self):
        paragraphs = [paragraph.text for paragraph in self._document().paragraphs]
        self.assertEqual(paragraphs[0], "知识点 1：窗口结构。【对口】【了解】")
        self.assertEqual(paragraphs[1], "CorelDraw窗口界面的组成（）【一般】【单选】")
        self.assertEqual(paragraphs[2], "A.标题栏、菜单栏")
        self.assertEqual(paragraphs[5], "D.色盘")
        self.assertEqual(paragraphs[6], "答案：【A】")
        # 题型顺序：单→多→判
        self.assertEqual(paragraphs[7], "CorelDRAW可链接的软件有（）【困难】【多选】")
        self.assertEqual(paragraphs[12], "答案：【BC】")
        self.assertEqual(
            paragraphs[13], "图片插入后可通过拖动边角调整大小。【简单】【判断】"
        )
        self.assertEqual(paragraphs[16], "答案：【A】")

    def test_judge_options_are_pinned(self):
        paragraphs = [paragraph.text for paragraph in self._document().paragraphs]
        index = paragraphs.index(
            "图片插入后可通过拖动边角调整大小。【简单】【判断】"
        )
        self.assertEqual(paragraphs[index + 1], "A.正确")
        self.assertEqual(paragraphs[index + 2], "B.错误")

    def test_stem_keeps_auto_numbering_and_spacing(self):
        xml = self._xml("word/document.xml")
        self.assertIn("<w:numPr>", xml)
        self.assertIn('<w:numId w:val="1"/>', xml)
        self.assertIn('w:before="218"', xml)
        self.assertIn('w:after="62"', xml)
        self.assertIn('w:beforeLines="70"', xml)
        self.assertIn('w:afterLines="20"', xml)

    def test_option_indent_and_font(self):
        xml = self._xml("word/document.xml")
        self.assertIn('w:left="426"', xml)
        self.assertIn('w:ascii="宋体"', xml)

    def test_title_paragraph_is_not_numbered(self):
        section = self._xml("word/document.xml").split("</w:p>")[0]
        self.assertNotIn("numPr", section)

    def test_invalid_questions_block_export(self):
        bad = make_point()
        bad.questions = [
            Question(
                stem="x",
                question_type="判断",
                options=["正确", "错误"],
                answer="C",
                internal_id="1-1-P-1",
            )
        ]
        with self.assertRaises(ValueError):
            self.exporter.export_word(bad)


class ZipExportTests(unittest.TestCase):
    def setUp(self):
        self.exporter = EvaluationExporter(TEMPLATE_DIR_MANAGER)
        self.point = make_point()
        self.point.questions = make_questions()
        self.second = make_point(name="工具栏", order=2)
        self.second.questions = [
            Question(
                stem="工具栏按钮（）",
                question_type="单选",
                difficulty="简单",
                options=["a", "b", "c", "d"],
                answer="A",
                internal_id="1-2-X-1",
            )
        ]
        self.draft = CourseDraft(
            course_name="CorelDraw",
            export_word=False,
            export_excel=True,
            chapters=[Chapter(name="界面介绍", knowledge_points=[self.point, self.second])],
        )
        self.draft.renumber()

    def test_zip_name_and_entries(self):
        name, data = self.exporter.export_zip(self.draft, date_text="20260920")
        self.assertEqual(name, "CorelDraw_评估题库_20260920.zip")

        entries = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual(
            entries,
            [
                "CorelDraw_1_界面介绍_1_窗口结构.xlsx",
                "CorelDraw_1_界面介绍_2_工具栏.xlsx",
            ],
        )

    def test_explicit_both_formats_override_draft_flags(self):
        _, data = self.exporter.export_zip(
            self.draft, date_text="20260920", want_word=True, want_excel=True
        )
        entries = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual(len(entries), 4)

    def test_word_and_excel_can_be_exported_together(self):
        draft = CourseDraft(
            course_name="CorelDraw",
            chapters=[Chapter(name="界面介绍", knowledge_points=[self.point])],
        )
        draft.renumber()
        _, data = self.exporter.export_zip(draft, date_text="20260920")
        entries = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual(
            entries,
            [
                "CorelDraw_1_界面介绍_1_窗口结构.docx",
                "CorelDraw_1_界面介绍_1_窗口结构.xlsx",
            ],
        )

    def test_duplicate_names_are_deduped(self):
        # 正常情况下 章节序/知识点序 保证文件名唯一；去重是防御性耍底，
        # 故直接验证 helper 的同名处理。
        from evaluation.exporter import _dedupe_name

        used: dict[str, int] = {}
        first = _dedupe_name("CorelDraw_1_A_1_B.xlsx", used)
        second = _dedupe_name("CorelDraw_1_A_1_B.xlsx", used)
        third = _dedupe_name("CorelDraw_1_A_1_B.xlsx", used)
        self.assertEqual(first, "CorelDraw_1_A_1_B.xlsx")
        self.assertEqual(second, "CorelDraw_1_A_1_B_2.xlsx")
        self.assertEqual(third, "CorelDraw_1_A_1_B_3.xlsx")

        # 大小写不敏感（Windows 文件名不区分大小写）
        self.assertEqual(
            _dedupe_name("coreldraw_1_a_1_b.xlsx", used), "coreldraw_1_a_1_b_4.xlsx"
        )

    def test_zip_entries_are_unique_for_two_chapters(self):
        second = make_point(name="窗口结构", order=1)
        second.questions = [
            Question(
                stem="重复题（）",
                question_type="单选",
                difficulty="简单",
                options=["a", "b", "c", "d"],
                answer="A",
                internal_id="2-1-X-1",
            )
        ]
        draft = CourseDraft(
            course_name="CorelDraw",
            export_word=False,
            chapters=[
                Chapter(name="界面介绍", knowledge_points=[self.point]),
                Chapter(name="界面介绍", knowledge_points=[second]),
            ],
        )
        draft.renumber()
        _, data = self.exporter.export_zip(draft, date_text="20260920")
        entries = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual(
            entries,
            [
                "CorelDraw_1_界面介绍_1_窗口结构.xlsx",
                "CorelDraw_2_界面介绍_1_窗口结构.xlsx",
            ],
        )

    def test_skipped_points_are_excluded(self):
        self.second.skipped = True
        _, data = self.exporter.export_zip(self.draft, date_text="20260920")
        entries = zipfile.ZipFile(io.BytesIO(data)).namelist()
        self.assertEqual(entries, ["CorelDraw_1_界面介绍_1_窗口结构.xlsx"])

    def test_write_exports_persists_the_zip(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            path, data = write_exports(
                self.exporter, self.draft, out_dir=tmp, date_text="20260920"
            )
            self.assertTrue(path.is_file())
            self.assertEqual(path.read_bytes(), data)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class DraftStoreFlushVolumeTests(unittest.TestCase):
    """§H-4：防抖确保生成过程中不会「每变一次就写盘」。"""

    def test_fifty_dirty_marks_produce_few_flushes(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            clock = [0.0]
            store = DraftStore(
                tmp, min_interval=2.0, every_n_writes=5, time_fn=lambda: clock[0]
            )
            draft = CourseDraft(course_name="X")

            for index in range(50):
                store.mark_dirty()
                # 模拟逐题生成：时间只推进 0.1s，主要依赖计数阈值
                clock[0] += 0.1
                store.flush_if_due(draft)

            # 每 5 次变更最多 1 次落盘，50 次 <= 10 次（实际更少）
            self.assertLessEqual(store.flush_count, 10)
            self.assertLess(store.flush_count, 50)

            # 强制落盘必须总是生效
            store.mark_dirty()
            self.assertTrue(store.flush(draft, force=True))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ============================================================== ai_client
class _StubResponse:
    def __init__(self, payload=None, *, headers=None, lines=None, json_error=False):
        self._payload = payload
        self.headers = headers or {"Content-Type": "application/json"}
        self._lines = lines or []
        self._json_error = json_error

    def raise_for_status(self):
        pass

    def json(self):
        if self._json_error:
            raise ValueError("not json")
        return self._payload

    def iter_lines(self, decode_unicode=False):
        return iter(self._lines)


class _StubSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls: list = []

    def post(self, url, **kwargs):
        self.calls.append((url, kwargs))
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response


def _ai_payload(questions: list) -> dict:
    return {"choices": [{"message": {"content": json.dumps({"questions": questions})}}]}


def _sample_question() -> dict:
    return {
        "stem": "题（）",
        "type": "单选",
        "difficulty": "简单",
        "options": ["a", "b", "c", "d"],
        "answer": "A",
        "confidence": 0.9,
        "rationale": "r",
    }


class AIClientTests(unittest.TestCase):
    def _request(self):
        return build_generation_request(
            course_name="CorelDraw",
            knowledge_point="窗口结构",
            tags=["标题栏"],
            question_count=1,
            type_counts={"单选": 1, "多选": 0, "判断": 0},
            difficulty_ratio={"简单": 100, "一般": 0, "困难": 0},
        )

    # ---- StructuredAIClient -----------------------------------------
    def test_generate_sends_openai_compatible_request(self):
        session = _StubSession([_StubResponse(_ai_payload([_sample_question()]))])
        client = StructuredAIClient(
            {"api_key": "sk-1", "base_url": "https://api.deepseek.com/v1", "model": "deepseek-chat"},
            session=session,
        )
        result = client.generate(self._request())

        self.assertEqual(len(result["questions"]), 1)
        url, kwargs = session.calls[0]
        self.assertEqual(url, "https://api.deepseek.com/v1/chat/completions")
        self.assertEqual(kwargs["json"]["response_format"], {"type": "json_object"})
        self.assertEqual(kwargs["json"]["model"], "deepseek-chat")
        self.assertEqual(kwargs["headers"]["Authorization"], "Bearer sk-1")

    def test_generate_retries_bad_json_then_succeeds(self):
        session = _StubSession(
            [
                _StubResponse({"choices": [{"message": {"content": "not json"}}]}),
                _StubResponse(_ai_payload([_sample_question()])),
            ]
        )
        client = StructuredAIClient({"api_key": "k"}, session=session, logger=lambda _: None)
        self.assertEqual(len(client.generate(self._request())["questions"]), 1)
        self.assertEqual(len(session.calls), 2)

    def test_generate_raises_after_exhausting_retries(self):
        session = _StubSession(
            [
                _StubResponse({"choices": [{"message": {"content": "{}"}}]}),
                _StubResponse({"choices": [{"message": {"content": "{}"}}]}),
                _StubResponse({"choices": [{"message": {"content": "{}"}}]}),
            ]
        )
        client = StructuredAIClient({"api_key": "k"}, session=session, logger=lambda _: None)
        with self.assertRaises(RuntimeError):
            client.generate(self._request(), retries=2)

    def test_missing_api_key_raises_config_error(self):
        client = StructuredAIClient({}, session=_StubSession([]))
        with self.assertRaises(AIConfigError):
            client.generate(self._request())
        with self.assertRaises(AIConfigError):
            client.complete_json(system="s", user="u")

    def test_complete_json_is_generic(self):
        session = _StubSession(
            [_StubResponse({"choices": [{"message": {"content": '{"tags": ["a"]}'}}]})]
        )
        client = StructuredAIClient({"api_key": "k"}, session=session)
        self.assertEqual(client.complete_json(system="s", user="u"), {"tags": ["a"]})

    def test_defaults_are_applied(self):
        client = StructuredAIClient({"api_key": "k"})
        self.assertEqual(client.base_url, "https://api.deepseek.com/v1")
        self.assertEqual(client.model, "deepseek-chat")

    # ---- Exa MCP ----------------------------------------------------
    def _search_client(self, session, config=None):
        from evaluation.ai_client import HttpMcpSearchClient

        return HttpMcpSearchClient(
            config or {"exa_mcp_url": "https://mcp.exa.ai/mcp"},
            session=session,
            logger=lambda _: None,
        )

    def test_exa_keyless_sse_response_is_parsed(self):
        inner = json.dumps(
            {"results": [{"title": "T1", "text": "snip1"}]}, ensure_ascii=False
        )
        envelope = {
            "jsonrpc": "2.0",
            "id": 1,
            "result": {"content": [{"type": "text", "text": inner}]},
        }
        session = _StubSession(
            [
                _StubResponse(
                    headers={"Content-Type": "text/event-stream; charset=utf-8"},
                    lines=["event: message", f"data: {json.dumps(envelope)}", ""],
                )
            ]
        )
        client = self._search_client(session)
        result = client.search("coreldraw 界面")

        self.assertIn("T1", result)
        self.assertIn("snip1", result)
        self.assertTrue(session.calls[0][1]["stream"])
        self.assertIn("text/event-stream", session.calls[0][1]["headers"]["Accept"])

    def test_exa_plain_json_response_is_parsed(self):
        session = _StubSession(
            [
                _StubResponse(
                    {"result": {"content": [{"type": "text", "text": "纯文本结果"}]}}
                )
            ]
        )
        self.assertIn("纯文本结果", self._search_client(session).search("q"))

    def test_exa_falls_back_to_api_key_when_keyless_fails(self):
        session = _StubSession(
            [
                RuntimeError("429 rate limited"),
                _StubResponse({"result": {"content": [{"type": "text", "text": "ok"}]}}),
            ]
        )
        client = self._search_client(
            session, {"exa_mcp_url": "https://mcp.exa.ai/mcp", "exa_api_key": "ek"}
        )
        self.assertEqual(client.search("q"), "ok")

        urls = [url for url, _ in session.calls]
        self.assertEqual(urls[0], "https://mcp.exa.ai/mcp")
        self.assertIn("exaApiKey=ek", urls[1])

    def test_exa_degrades_silently_when_everything_fails(self):
        from evaluation.ai_client import HttpMcpSearchClient

        session = _StubSession([RuntimeError("boom"), RuntimeError("boom2")])
        client = HttpMcpSearchClient(
            {"exa_mcp_url": "https://mcp.exa.ai/mcp", "exa_api_key": "ek"},
            session=session,
            logger=lambda _: None,
        )
        # 检索失败绝不能抛异常（否则会中断出题）
        self.assertEqual(client.search("q"), "")

    def test_exa_missing_url_returns_empty(self):
        self.assertEqual(self._search_client(_StubSession([]), {"exa_mcp_url": ""}).search("q"), "")
        self.assertEqual(self._search_client(_StubSession([])).search("   "), "")

    # ---- 解析工具 ---------------------------------------------------
    def test_parse_json_content_variants(self):
        from evaluation.ai_client import parse_json_content

        self.assertEqual(parse_json_content('{"a": 1}'), {"a": 1})
        self.assertEqual(parse_json_content('```json\n{"a": 2}\n```'), {"a": 2})
        self.assertEqual(parse_json_content('blah {"a": 3} tail'), {"a": 3})
        self.assertIsNone(parse_json_content("nope"))
        self.assertIsNone(parse_json_content(""))
        self.assertEqual(parse_json_content({"a": 4}), {"a": 4})

    def test_should_search(self):
        from evaluation.ai_client import should_search

        self.assertTrue(should_search(True, True))
        self.assertFalse(should_search(True, False))
        self.assertFalse(should_search(False, True))


class RealModelOutputTests(unittest.TestCase):
    """用**真实抓到的** AI 返回做回归（deepseek-v4-flash）。

    事故：``response_format={'type':'json_object'}`` 只保证「是合法 JSON」，
    **不保证字段名**；当时请求里根本没带字段契约（QUESTION_SCHEMA 构造了
    却从未发出），且解析只认 ``stem``，导致模型返回的 ``question`` 被丢掉，
    30 道题全部变成空题干。
    """

    # 真实抓取的原始响应（字段名 'question' 是模型自创的）
    REAL_PAYLOAD = {
        "questions": [
            {
                "type": "单选",
                "question": "下列关于CorelDRAW窗口界面的说法，正确的是（）",
                "options": [
                    "菜单栏、工具箱、属性栏都属于CorelDRAW窗口的组成部分",
                    "属性栏固定不变，不随所选工具或对象而变化",
                    "标准工具栏中只包含绘图工具",
                    "状态栏用于设置当前工具的详细参数",
                ],
                "answer": "A",
                "difficulty": "简单",
                "confidence": 0.95,
                "rationale": "考查CorelDRAW窗口组成及工具栏、属性栏的基本概念与作用。",
            }
        ]
    }

    def _controller(self, records=None):
        from evaluation.ai_client import StructuredAIClient
        from evaluation.generator import GenerationController

        client = StructuredAIClient({"api_key": "k"}, session=_StubSession([]))
        queue = ClarificationQueue(
            sensitivity="标准", on_record=(records.append if records else None)
        )
        logs: list[str] = []
        return GenerationController(client, None, queue, logger=logs.append), logs

    def test_contract_asks_for_spread_out_answers(self):
        """真实跑出的 10 题里 3 道单选答案全是 A，学生可靠猜「选 A」得分；
        因此字段契约必须要求打散答案位置。"""
        request = build_generation_request(
            course_name="CorelDraw",
            knowledge_point="窗口界面",
            tags=["窗口组成"],
            question_count=3,
            type_counts={"单选": 3, "多选": 0, "判断": 0},
            difficulty_ratio={"简单": 100, "一般": 0, "困难": 0},
        )
        text = request["system"]
        self.assertIn("答案位置", text)
        self.assertIn("不能集中在 A", text)

    def test_answer_distribution_warning_detects_all_a(self):
        from evaluation.quality import answer_distribution_warning

        questions = [
            Question(stem=f"q{i}", question_type="单选", answer="A") for i in range(5)
        ]
        warnings = answer_distribution_warning(questions)
        self.assertTrue(warnings)
        self.assertIn("A", warnings[0])

        spread = [
            Question(stem=f"q{i}", question_type="单选", answer="ABCD"[i % 4])
            for i in range(8)
        ]
        self.assertEqual(answer_distribution_warning(spread), [])

    def test_stem_alias_is_picked_up_and_warned(self):
        controller, logs = self._controller()
        point = make_point()
        questions = controller._normalize_payload(
            self.REAL_PAYLOAD, point, {"单选": 1}, CourseDraft(), on_clarification=None
        )

        self.assertEqual(len(questions), 1)
        question = questions[0]
        self.assertEqual(
            question.stem, "下列关于CorelDRAW窗口界面的说法，正确的是（）"
        )
        self.assertEqual(question.answer, "A")
        self.assertEqual(question.question_type, "单选")
        self.assertEqual(question.difficulty, "简单")
        self.assertEqual(len(question.options), 4)
        self.assertEqual(validate_question(question), [])
        self.assertTrue(
            any("非标准字段名" in line for line in logs),
            "应记录非标准字段名告警，避免静默丢题干",
        )

    def test_prompt_carries_explicit_field_contract(self):
        """字段契约必须真的进到请求里（曾经只存在 request['response_schema'] 里没发出）。"""
        request = build_generation_request(
            course_name="CorelDraw",
            knowledge_point="窗口界面",
            tags=["窗口组成"],
            question_count=1,
            type_counts={"单选": 1, "多选": 0, "判断": 0},
            difficulty_ratio={"简单": 100, "一般": 0, "困难": 0},
        )
        system = request["system"]
        user = request["user"]

        for key in (
            "stem",
            "type",
            "difficulty",
            "options",
            "answer",
            "confidence",
            "rationale",
        ):
            self.assertIn(f'"{key}"', system, f"字段 {key} 必须出现在契约里")

        self.assertIn("stem", system)
        self.assertIn("不是 question", system)
        self.assertIn("字段名", user)

    def test_pick_field_alias_table(self):
        from evaluation.ai_client import pick_field

        raw = {"question": "Q", "choices": ["a", "b"], "correctAnswer": "A"}
        self.assertEqual(pick_field(raw, "stem"), ("Q", "question"))
        self.assertEqual(pick_field(raw, "options")[0], ["a", "b"])
        self.assertEqual(pick_field(raw, "answer"), ("A", "correctAnswer"))
        self.assertEqual(pick_field(raw, "missing", "fallback"), ("fallback", None))
        # 空值不算命中，避免 ``"stem": ""`` 盖住后面的有效别名
        self.assertEqual(pick_field({"stem": "", "question": "Q"}, "stem"), ("Q", "question"))

    def test_judge_options_still_pinned_with_aliases(self):
        controller, _ = self._controller()
        payload = {
            "questions": [
                {
                    "question": "按F9可放大视图。（）",
                    "type": "判断",
                    "difficulty": "简单",
                    "choices": ["对", "错"],
                    "answer": "A",
                    "confidence": 0.9,
                    "rationale": "r",
                }
            ]
        }
        questions = controller._normalize_payload(
            payload, make_point(), {"判断": 1}, CourseDraft(), on_clarification=None
        )
        self.assertEqual(questions[0].options, ["正确", "错误"])
        self.assertEqual(validate_question(questions[0]), [])

    def test_answer_list_is_joined(self):
        controller, _ = self._controller()
        payload = {
            "questions": [
                {
                    "question": "可导出为哪些格式（）",
                    "type": "多选",
                    "difficulty": "一般",
                    "options": ["JPG", "PSD", "AI", "MP4"],
                    "answer": ["C", "A", "B"],
                    "confidence": 0.8,
                    "rationale": "r",
                }
            ]
        }
        questions = controller._normalize_payload(
            payload, make_point(), {"多选": 1}, CourseDraft(), on_clarification=None
        )
        self.assertEqual(questions[0].answer, "ABC")
        self.assertEqual(validate_question(questions[0]), [])


TEMPLATE_DIR_MANAGER = TemplateManager(PLUGIN_DIR)


if __name__ == "__main__":
    unittest.main()
