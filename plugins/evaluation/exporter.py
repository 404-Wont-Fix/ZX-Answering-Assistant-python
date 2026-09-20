"""
平台模板导出器 - Platform-template exporters.

设计要点（对应 PRD §10 修订 + §H 防御性设计）：

- **导出粒度 = 知识点**：一个知识点一个 Word + 一个 Excel；
- Excel 以 ``kp-question-excel.xlsx`` 为基底：清掉模板示例行、写 R5 起数据行，
  **合并单元格防御性写入**（§H-3），不写示例题、不写「删除示例」提示行；
- Word 以 ``kp-import-template.docx`` 为基底：**节点级 XML 克隆**（§H-2）
  保住自动编号 / 缩进 / 宋体 rFonts / 字号，页眉标题改写、页脚 PAGE/NUMPAGES 域原样保留；
- 文件名清洗防 ``OSError``（§H-6）。
"""

from __future__ import annotations

import copy
import io
import re
import zipfile
from pathlib import Path
from typing import Iterable, Sequence

from .models import (
    JUDGE_OPTIONS,
    MAX_OPTIONS,
    OPTION_LETTERS,
    CourseDraft,
    KnowledgePoint,
    Question,
)
from .template_manager import TemplateManager
from .validators import sort_questions, validate_questions

# ---------------------------------------------------------------- 常量
EXCEL_SHEET = "Sheet1"
EXCEL_NOTE_ROW = 1
EXCEL_META_ROW = 2
EXCEL_META_VALUE_ROW = 3
EXCEL_HEADER_ROW = 4
EXCEL_FIRST_DATA_ROW = 5

# 合并区（row=3 上的 min_col → max_col）：A3:D3 / E3:H3 / I3:L3
EXCEL_MERGES: dict[int, int] = {1: 4, 5: 8, 9: 12}

WORD_HEADER_SUFFIX = "”知识点评估题"
WORD_TITLE_PREFIX = "知识点 "

# 文件名清洗：Windows 非法字符
_ILLEGAL_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]+')
_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}
_MAX_SEGMENT_LENGTH = 80


# 判断题结尾的作答位括号：模板样例的判断题是陈述句，不带 （）
_TRAILING_BRACKETS = re.compile(r"[（(]\s*[）)]\s*$")


def normalize_export_stem(question: Question) -> str:
    """
    导出用题干（**Word 与 Excel 共用**，避免两份产物不一致）。

    只做一件事：判断题去掉模型可能多写的尾部作答位括号 —— 判断题是陈述句，
    模板的 Word/Excel 判断题样例均不带 （）:

    - Word 样例：``图片插入后，你可以通过拖动图片的边角来调整其大小。``
    - Excel 样例：``图片插入后，你可以通过拖动图片的边角来调整其大小。``
    """
    stem = str(question.stem or "").strip()
    if question.question_type == "判断":
        stem = _TRAILING_BRACKETS.sub("", stem).strip()
    return stem


def format_question_stem(question: Question) -> str:
    """
    题干行的平台格式（与模板逐行一致）：``{题干}（）【{难度}】【{题型}】``。

    - 单选/多选：无括号时补 ``（）`` 作答位；
    - 判断：不加作答位括号（陈述句不该有）。
    """
    stem = normalize_export_stem(question)

    if question.question_type != "判断" and "（" not in stem and "(" not in stem:
        stem = f"{stem}（）"

    return f"{stem}【{question.difficulty}】【{question.question_type}】"


def safe_filename_segment(name: str, *, fallback: str = "evaluation") -> str:
    """
    清洗单个文件名片段（§H-6）。

    - 替换 Windows 非法字符为 ``_``；
    - 去掉首尾空白与点；
    - 折叠连续下划线；
    - 规避 Windows 保留名（前置 ``_``）；
    - 超长截断到 80 字符。
    """
    text = _ILLEGAL_FILENAME_CHARS.sub("_", str(name or ""))
    text = text.strip().strip(".")
    text = re.sub(r"_{2,}", "_", text).strip("_")

    if not text:
        return fallback

    if text.upper() in _RESERVED_NAMES:
        text = f"_{text}"

    if len(text) > _MAX_SEGMENT_LENGTH:
        text = text[:_MAX_SEGMENT_LENGTH].rstrip("_.")

    return text or fallback


# 向后兼容别名（计划里提到 _safe_stem）
_safe_stem = safe_filename_segment


def knowledge_point_file_stem(
    course_name: str,
    point: KnowledgePoint,
) -> str:
    """知识点文件名（不含扩展名）：``{课程名}_{章节序号}_{章节名}_{知识点序号}_{知识点名}``。"""
    return "_".join(
        [
            safe_filename_segment(course_name, fallback="课程"),
            str(point.chapter_order),
            safe_filename_segment(point.chapter_name, fallback="章节"),
            str(point.order),
            safe_filename_segment(point.name, fallback="知识点"),
        ]
    )


def zip_name(course_name: str, date_text: str) -> str:
    """zip 包名：``{课程名}_评估题库_{日期}.zip``。"""
    course = safe_filename_segment(course_name, fallback="课程")
    return f"{course}_评估题库_{safe_filename_segment(date_text, fallback='日期')}.zip"


class EvaluationExporter:
    """
    模板驱动的导出器。

    Args:
        templates: 已初始化的 :class:`TemplateManager`（或插件目录路径）
    """

    def __init__(self, templates: TemplateManager | str | Path) -> None:
        self.templates = (
            templates
            if isinstance(templates, TemplateManager)
            else TemplateManager(templates)
        )

    # ================================================================ Excel
    def export_excel(self, point: KnowledgePoint) -> bytes:
        """
        导出单个知识点的 Excel（平台可直接导入）。

        Raises:
            ValueError: 题目校验不通过（FR6.9，消息已定位到题）
        """
        questions = sort_questions(point.questions)
        problems = validate_questions(questions)
        if problems:
            raise ValueError("题目校验失败: " + "；".join(problems))

        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - openpyxl 是根依赖
            raise RuntimeError("缺少 openpyxl，无法导出 Excel") from exc

        workbook = load_workbook(str(self.templates.get_path("excel")))

        try:
            sheet = (
                workbook[EXCEL_SHEET]
                if EXCEL_SHEET in workbook.sheetnames
                else workbook.worksheets[0]
            )

            # 元信息（合并区防御性写入）
            self._write_merged(sheet, EXCEL_META_VALUE_ROW, 1, point.name)
            self._write_merged(sheet, EXCEL_META_VALUE_ROW, 5, point.level)
            self._write_merged(sheet, EXCEL_META_VALUE_ROW, 9, point.knowledge_type)

            # 清掉模板自带的示例题 + 「示例/请删除」提示行
            if sheet.max_row >= EXCEL_FIRST_DATA_ROW:
                sheet.delete_rows(
                    EXCEL_FIRST_DATA_ROW,
                    sheet.max_row - EXCEL_FIRST_DATA_ROW + 1,
                )

            # 记录数据行样式原型（保留模板的宋体 11pt）
            prototype = self._capture_row_styles(sheet, EXCEL_FIRST_DATA_ROW - 1)

            for offset, question in enumerate(questions):
                row_index = EXCEL_FIRST_DATA_ROW + offset
                self._write_data_row(sheet, row_index, offset + 1, question, prototype)

            self._assert_excel_structure(sheet, expected_rows=len(questions))

            buffer = io.BytesIO()
            workbook.save(buffer)
        finally:
            workbook.close()

        return buffer.getvalue()

    @staticmethod
    def _capture_row_styles(sheet, reference_row: int) -> dict[int, object]:
        """抓取参考行的单元格样式，供新数据行复用（保持模板字体）。"""
        styles: dict[int, object] = {}
        for column in range(1, 5 + MAX_OPTIONS):
            try:
                styles[column] = copy.copy(sheet.cell(row=reference_row, column=column)._style)
            except (AttributeError, IndexError):
                continue
        return styles

    def _write_data_row(
        self,
        sheet,
        row_index: int,
        serial: int,
        question: Question,
        prototype: dict[int, object],
    ) -> None:
        values: list[object] = [
            serial,
            normalize_export_stem(question),
            question.difficulty,
            question.question_type,
            question.answer,
        ]
        if question.question_type == "判断":
            options = list(JUDGE_OPTIONS)
        else:
            options = [str(item) for item in (question.options or [])][:MAX_OPTIONS]
        values.extend(options)

        for index, value in enumerate(values):
            column = index + 1
            cell = sheet.cell(row=row_index, column=column)
            cell.value = value
            style = prototype.get(column)
            if style is not None:
                cell._style = style

    @staticmethod
    def _write_merged(sheet, row: int, min_col: int, value) -> None:
        """
        向合并区左上角写值，并**先确保合并状态正确**（§H-3）。

        WPS 另存等因素可能导致合并范围偏移或丢失，因此：

        - 若与目标范围不一致 → 先 ``unmerge`` 再 ``merge``；
        - 若完全未合并 → 直接 ``merge``；
        - 最后才写左上角单元格的值。
        """
        max_col = EXCEL_MERGES.get(min_col, min_col)

        existing = None
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row <= row <= merged.max_row and merged.min_col <= min_col <= merged.max_col:
                existing = merged
                break

        mismatched = existing is not None and (
            existing.min_row != row
            or existing.max_row != row
            or existing.min_col != min_col
            or existing.max_col != max_col
        )
        if mismatched:
            sheet.unmerge_cells(str(existing))
            existing = None

        if existing is None:
            sheet.merge_cells(
                start_row=row,
                start_column=min_col,
                end_row=row,
                end_column=max_col,
            )

        sheet.cell(row=row, column=min_col).value = value

    @staticmethod
    def _assert_excel_structure(sheet, *, expected_rows: int) -> None:
        """导出后自检：表头未被破坏、数据从 R5 开始（防静默错位）。"""
        headers = [
            sheet.cell(row=EXCEL_HEADER_ROW, column=column).value
            for column in range(1, 14)
        ]
        expected_prefixes = [
            "序号",
            "题目",
            "难度",
            "类型",
            "正确答案项",
            *[f"答案{letter}" for letter in OPTION_LETTERS],
        ]

        for index, prefix in enumerate(expected_prefixes):
            value = str(headers[index] or "")
            if not value.startswith(prefix):
                raise ValueError(
                    f"Excel 表头异常：R{EXCEL_HEADER_ROW}C{index + 1} 期望以「{prefix}」开头，"
                    f"实际为「{value}」"
                )

        if expected_rows and sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=1).value != 1:
            raise ValueError(
                f"Excel 数据行异常：R{EXCEL_FIRST_DATA_ROW} 序号应为 1，"
                f"实际为「{sheet.cell(row=EXCEL_FIRST_DATA_ROW, column=1).value}」"
            )

    # ================================================================ Word
    def export_word(self, point: KnowledgePoint) -> bytes:
        """
        导出单个知识点的 Word（与平台模板逐行一致）。

        Raises:
            RuntimeError: 缺少 python-docx
            ValueError: 题目校验不通过
        """
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，请先安装插件依赖") from exc

        questions = sort_questions(point.questions)
        problems = validate_questions(questions)
        if problems:
            raise ValueError("题目校验失败: " + "；".join(problems))

        document = docx.Document(str(self.templates.get_path("word")))

        self._set_header_title(document, point.name)
        prototypes = self._capture_paragraph_prototypes(document)
        self._rewrite_title_paragraph(document, point, questions)
        self._clear_body_paragraphs(document)
        self._append_question_paragraphs(document, questions, prototypes)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    # ---------------------------------------------------------- Word 辅助
    @staticmethod
    def _set_header_title(document, point_name: str) -> None:
        """
        改写页眉为 ``“{知识点名}”知识点评估题``。

        只替换已有 run 的文本（不新增 run、不动 rId），保持加粗与字号。
        """
        for section in document.sections:
            for paragraph in section.header.paragraphs:
                runs = paragraph.runs
                if not runs:
                    continue
                if not any(WORD_HEADER_SUFFIX in run.text for run in runs):
                    continue

                suffix_index = next(
                    index
                    for index, run in enumerate(runs)
                    if WORD_HEADER_SUFFIX in run.text
                )
                # 名称所在的 run：suffix 之前、不是纯引号/制表符的那个
                target_index = None
                for index in range(suffix_index - 1, -1, -1):
                    stripped = runs[index].text.strip()
                    if stripped and stripped not in ("“", "”", '"'):
                        target_index = index
                        break

                if target_index is None:
                    # 退路：把名称塞进 suffix 前一个 run，并补齐引号
                    runs[max(suffix_index - 1, 0)].text = f"“{point_name}”"
                    runs[suffix_index].text = "知识点评估题"
                    continue

                runs[target_index].text = point_name
                return

    @staticmethod
    def _capture_paragraph_prototypes(document) -> dict[str, object]:
        """
        从模板正文抓取「题干 / 选项 / 答案」三种段落的 XML 原型（§H-2）。

        注意必须取 ``_element``（XML 节点）而不是 ``Paragraph`` 包装对象，
        否则克隆后可能缺少 ``pPr``（自动编号/缩进/间距）与 ``rPr``（宋体/字号）。
        """
        from docx.oxml.ns import qn  # type: ignore

        prototypes: dict[str, object] = {}
        paragraphs = document.paragraphs

        for paragraph in paragraphs[1:]:
            element = paragraph._element
            if element.find(qn("w:pPr")) is None:
                continue
            if element.find(f".//{qn('w:numPr')}") is not None and "stem" not in prototypes:
                prototypes["stem"] = copy.deepcopy(element)
                continue

            text = paragraph.text.strip()
            if text.startswith("答案") and "answer" not in prototypes:
                prototypes["answer"] = copy.deepcopy(element)
                continue
            if (
                len(text) > 1
                and text[0] in OPTION_LETTERS
                and text[1] in (".", "．")
                and "option" not in prototypes
            ):
                prototypes["option"] = copy.deepcopy(element)

        if "stem" not in prototypes and len(paragraphs) > 1:
            prototypes["stem"] = copy.deepcopy(paragraphs[1]._element)
        if "option" not in prototypes and len(paragraphs) > 2:
            prototypes["option"] = copy.deepcopy(paragraphs[2]._element)
        if "answer" not in prototypes and len(paragraphs) > 3:
            prototypes["answer"] = copy.deepcopy(paragraphs[3]._element)

        return prototypes

    @staticmethod
    def _rewrite_title_paragraph(document, point: KnowledgePoint, questions: Sequence[Question]) -> None:
        """改写正文首段：``知识点 {序号}：{名称}。【{等级}】【{类型}】``。"""
        paragraphs = document.paragraphs
        if not paragraphs:
            return

        title = (
            f"{WORD_TITLE_PREFIX}{point.order}：{point.name}。"
            f"【{point.level}】【{point.knowledge_type}】"
        )
        paragraph = paragraphs[0]
        runs = paragraph.runs
        if not runs:
            paragraph.add_run(title)
            return

        runs[0].text = title
        for run in runs[1:]:
            run.text = ""

    @staticmethod
    def _clear_body_paragraphs(document) -> None:
        """删除正文中除首段（标题）以外的所有段落；``sectPr`` 保留。"""
        for paragraph in document.paragraphs[1:]:
            element = paragraph._element
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)

    def _append_question_paragraphs(
        self,
        document,
        questions: Sequence[Question],
        prototypes: dict[str, object],
    ) -> None:
        """按导出顺序克隆段落并写入题目内容。"""
        from docx.oxml.ns import qn  # type: ignore

        body = document.element.body
        sect_pr = body.find(qn("w:sectPr"))

        def emit(kind: str, text: str) -> None:
            # 注意：lxml 元素不能裸用于布尔判断（空元素为假，且未来版本恒为真）
            prototype = prototypes.get(kind)
            if prototype is None:
                prototype = prototypes.get("stem")
            if prototype is None:
                return
            element = copy.deepcopy(prototype)
            _set_paragraph_text(element, text)
            if sect_pr is not None:
                sect_pr.addprevious(element)
            else:
                body.append(element)

        for question in questions:
            emit("stem", format_question_stem(question))

            if question.question_type == "判断":
                options = list(JUDGE_OPTIONS)
            else:
                options = [str(item) for item in (question.options or [])][:MAX_OPTIONS]

            for index, option in enumerate(options):
                letter = OPTION_LETTERS[index] if index < len(OPTION_LETTERS) else "?"
                emit("option", f"{letter}.{option}")

            emit("answer", f"答案：【{question.answer}】")

    # ================================================================ 批量/ZIP
    def export_point(
        self,
        point: KnowledgePoint,
        *,
        want_word: bool | None = None,
        want_excel: bool | None = None,
    ) -> dict[str, bytes]:
        """导出单个知识点，返回 ``{扩展名: 内容}``。"""
        result: dict[str, bytes] = {}
        if want_word is None or want_word:
            result[".docx"] = self.export_word(point)
        if want_excel is None or want_excel:
            result[".xlsx"] = self.export_excel(point)
        return result

    def export_zip(
        self,
        draft: CourseDraft,
        *,
        date_text: str,
        want_word: bool | None = None,
        want_excel: bool | None = None,
    ) -> tuple[str, bytes]:
        """
        打包导出：包内按 章节序 → 知识点序 排列，文件名按 §6 固定规则。

        ``want_word`` / ``want_excel`` 为 None 时跟随 ``draft.export_word`` /
        ``draft.export_excel``（Step6 的导出格式勾选）。

        Returns:
            (zip 文件名, zip 内容字节)
        """
        include_word = draft.export_word if want_word is None else bool(want_word)
        include_excel = draft.export_excel if want_excel is None else bool(want_excel)
        if not include_word and not include_excel:
            raise ValueError("请至少选择一种导出格式")

        course = safe_filename_segment(draft.course_name, fallback="课程")
        buffer = io.BytesIO()
        used_names: dict[str, int] = {}

        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for chapter in draft.chapters:
                for point in chapter.knowledge_points:
                    if point.skipped or not point.questions:
                        continue

                    stem = knowledge_point_file_stem(course, point)
                    for suffix, content in sorted(
                        self.export_point(
                            point,
                            want_word=include_word,
                            want_excel=include_excel,
                        ).items()
                    ):
                        name = _dedupe_name(f"{stem}{suffix}", used_names)
                        archive.writestr(name, content)

        return zip_name(draft.course_name, date_text), buffer.getvalue()


def _dedupe_name(name: str, used: dict[str, int]) -> str:
    """同名 zip 条目自动加 ``_2`` / ``_3`` ...，避免静默覆盖。"""
    key = name.lower()
    count = used.get(key, 0)
    used[key] = count + 1
    if count == 0:
        return name

    stem, dot, suffix = name.rpartition(".")
    if not dot:
        return f"{name}_{count + 1}"
    return f"{stem}_{count + 1}.{suffix}"


def _set_paragraph_text(element, text: str) -> None:
    """
    把 XML 段落节点的文本替换为 ``text``（保留首个 run 的 rPr）。

    多余 run 一律清空文本而不是删除节点 —— 删除会连带影响
    ``w:rPr`` 继承与可能的书签/关系引用（§H-2）。
    """
    from docx.oxml.ns import qn  # type: ignore

    runs = element.findall(f".//{qn('w:r')}")
    if not runs:
        return

    first = runs[0]
    text_nodes = first.findall(qn("w:t"))
    if text_nodes:
        text_nodes[0].text = text
        text_nodes[0].set(qn("xml:space"), "preserve")
        for extra in text_nodes[1:]:
            extra.text = ""
    else:
        from docx.oxml import OxmlElement  # type: ignore

        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = text
        first.append(node)

    for run in runs[1:]:
        for node in run.findall(qn("w:t")):
            node.text = ""


def write_exports(
    exporter: EvaluationExporter,
    draft: CourseDraft,
    *,
    out_dir: str | Path,
    date_text: str,
    want_word: bool | None = None,
    want_excel: bool | None = None,
) -> tuple[Path, bytes]:
    """
    直接落盘 zip 到 ``out_dir``（§E 落盘约定：先写 output/evaluation/，再另存）。

    Returns:
        (zip 路径, zip 字节)
    """
    name, data = exporter.export_zip(
        draft, date_text=date_text, want_word=want_word, want_excel=want_excel
    )
    directory = Path(out_dir)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    path.write_bytes(data)
    return path, data


def exported_questions(point: KnowledgePoint) -> Iterable[Question]:
    """便捷方法：某知识点将实际被导出的题目（已排序）。"""
    return sort_questions(point.questions)
