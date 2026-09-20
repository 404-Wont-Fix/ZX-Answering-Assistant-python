"""
样例题库与知识点导入 - Sample-bank and knowledge-point import helpers.

对应 PRD：

- FR2.3 知识点支持 Excel 批量导入；
- FR2.4 ［AI 提取候选］从已上传考纲/大纲自动列出候选知识点；
- FR1.3 样例题库**仅用于对齐输出格式与难度风格，不作内容依据**。

考纲/大纲支持 ``.xlsx`` / ``.docx``（PDF 不支持，调用方据此提示）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .models import JUDGE_OPTIONS, MAX_OPTIONS, OPTION_LETTERS, Question

SUPPORTED_SYLLABUS_SUFFIXES = (".xlsx", ".docx")

# 知识点导入时可能出现的表头别名（取首个命中列）
_KNOWLEDGE_HEADER_CANDIDATES = (
    "知识点",
    "知识点名称",
    "knowledge",
    "knowledge point",
    "考点",
)


def import_knowledge_points_xlsx(path: str | Path) -> list[str]:
    """
    从 Excel 批量导入知识点名称。

    逐行读取第一个非空单元格文本；自动跳表头行；去重且保持原顺序。

    Args:
        path: xlsx 文件路径

    Returns:
        list[str]: 知识点名称列表（去重、去空白）
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl 是根依赖
        raise RuntimeError("缺少 openpyxl，无法导入知识点") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    names: list[str] = []
    seen: set[str] = set()

    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row or ():
                    text = str(cell).strip() if cell is not None else ""
                    if not text:
                        continue
                    if _looks_like_knowledge_header(text):
                        break
                    if text not in seen:
                        seen.add(text)
                        names.append(text)
                    # 一行只取第一个非空单元格
                    break
    finally:
        workbook.close()

    return names


def _looks_like_knowledge_header(text: str) -> bool:
    lowered = text.lower()
    return any(
        candidate in lowered or lowered in candidate
        for candidate in _KNOWLEDGE_HEADER_CANDIDATES
    )


def import_question_examples(path: str | Path) -> list[dict[str, Any]]:
    """
    导入样例题库（仅用于格式/难度对齐）。

    支持平台标准 Excel（``kp-question-excel.xlsx`` 结构）与 JSON。

    Returns:
        list[dict]: 每项含 ``stem / type / difficulty / options / answer``
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".json":
        return _import_json_examples(path)
    if suffix == ".xlsx":
        return _import_xlsx_examples(path)
    if suffix == ".docx":
        return _import_docx_examples(path)

    raise ValueError(f"不支持的样例题库格式：{suffix}")


# ---------------------------------------------------------------- Excel
def _import_xlsx_examples(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("缺少 openpyxl，无法导入样例题库") from exc

    workbook = load_workbook(path, data_only=True)
    examples: list[dict[str, Any]] = []

    try:
        for sheet in workbook.worksheets:
            header_row = _find_question_header_row(sheet)
            if header_row is None:
                continue

            columns = _map_question_columns(sheet, header_row)
            for row_index in range(header_row + 1, sheet.max_row + 1):
                stem = _cell_text(sheet, row_index, columns.get("stem"))
                if not stem:
                    continue

                question_type = _cell_text(sheet, row_index, columns.get("type")) or "单选"
                difficulty = _cell_text(sheet, row_index, columns.get("difficulty")) or "一般"
                answer = (_cell_text(sheet, row_index, columns.get("answer")) or "").upper()

                # 平台模板 R21 之类的「示例/请删除」提示行不是题目，必须跳过
                if _is_note_row(stem):
                    continue

                options: list[str] = []
                for letter_index, letter in enumerate(OPTION_LETTERS):
                    if question_type == "判断" and letter_index >= len(JUDGE_OPTIONS):
                        break
                    value = _cell_text(sheet, row_index, columns.get(f"option_{letter}"))
                    if value:
                        options.append(value)

                examples.append(
                    {
                        "stem": stem,
                        "type": question_type,
                        "difficulty": difficulty,
                        "options": options,
                        "answer": answer,
                    }
                )
    finally:
        workbook.close()

    return examples


# 平台模板里的「示例/请删除」提示行特征（如 B21「注意：以上三个试题为示例，正式录入的时候请删除」）
_NOTE_ROW_MARKERS = ("请删除", "正式录入", "示例，", "示例,")


def _is_note_row(text: str) -> bool:
    """判断某行是否是模板自带的说明/示例提示行（不是真题目）。"""
    stripped = (text or "").strip()
    if not stripped:
        return True
    if stripped.startswith("注意"):
        return True
    return any(marker in stripped for marker in _NOTE_ROW_MARKERS)


def _find_question_header_row(sheet) -> int | None:
    """找到含「题目」表头的行（平台模板为 R4）。"""
    for row_index in range(1, min(sheet.max_row, 20) + 1):
        for cell in sheet[row_index]:
            text = str(cell.value).strip() if cell.value is not None else ""
            if text in ("题目", "题干", "question"):
                return row_index
    return None


def _map_question_columns(sheet, header_row: int) -> dict[str, int]:
    """把表头文本映射到列号。"""
    mapping: dict[str, int] = {}
    for cell in sheet[header_row]:
        text = str(cell.value).strip() if cell.value is not None else ""
        if not text:
            continue

        if text in ("题目", "题干", "question"):
            mapping["stem"] = cell.column
        elif text.startswith("难度"):
            mapping["difficulty"] = cell.column
        elif text.startswith("类型"):
            mapping["type"] = cell.column
        elif text.startswith("正确答案"):
            mapping["answer"] = cell.column
        elif text.startswith("答案") and len(text) == 3:
            # 答案A / 答案B ...
            letter = text[-1].upper()
            if letter in OPTION_LETTERS:
                mapping[f"option_{letter}"] = cell.column
    return mapping


def _cell_text(sheet, row_index: int, column_index: int | None) -> str:
    if not column_index:
        return ""
    value = sheet.cell(row=row_index, column=column_index).value
    return str(value).strip() if value is not None else ""


# ---------------------------------------------------------------- JSON
def _import_json_examples(path: Path) -> list[dict[str, Any]]:
    import json

    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"样例题库 JSON 解析失败：{exc}") from exc

    # 兼容 {课程: {章节: {知识点: [题]}}} 与 [题, ...] 两种结构
    items = _flatten_json_questions(payload)
    examples: list[dict[str, Any]] = []

    for item in items:
        if not isinstance(item, dict):
            continue
        stem = str(item.get("question") or item.get("stem") or item.get("title") or "").strip()
        if not stem:
            continue
        examples.append(
            {
                "stem": stem,
                "type": _guess_type(item),
                "difficulty": "一般",
                "options": _extract_json_options(item),
                "answer": str(item.get("answer") or "").upper(),
            }
        )

    return examples


def _flatten_json_questions(payload: Any) -> list[Any]:
    """把各种嵌套结构里的题目列表摊平。"""
    if isinstance(payload, list):
        return payload
    if not isinstance(payload, dict):
        return []

    items: list[Any] = []
    for value in payload.values():
        if isinstance(value, list):
            items.extend(value)
        elif isinstance(value, dict):
            items.extend(_flatten_json_questions(value))
    return items


def _extract_json_options(item: dict[str, Any]) -> list[str]:
    options: list[str] = []
    for letter in OPTION_LETTERS[:MAX_OPTIONS]:
        value = item.get(letter) or item.get(letter.lower())
        if value:
            options.append(str(value))
    return options


def _guess_type(item: dict[str, Any]) -> str:
    raw = str(item.get("type") or item.get("question_type") or "").strip()
    if raw in ("单选", "多选", "判断"):
        return raw
    answer = str(item.get("answer") or "").upper()
    if len([ch for ch in answer if ch.isalpha()]) > 1:
        return "多选"
    return "单选"


# ---------------------------------------------------------------- 考纲
def extract_syllabus_text(path: str | Path) -> str:
    """
    从考纲/大纲（``.docx`` / ``.xlsx``）提取纯文本，供 AI 抽取候选知识点。

    Raises:
        ValueError: 不支持的格式（如 PDF）
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _import_docx_text(path)
    if suffix == ".xlsx":
        return _import_xlsx_text(path)

    raise ValueError(
        f"考纲仅支持 {SUPPORTED_SYLLABUS_SUFFIXES[0]} / {SUPPORTED_SYLLABUS_SUFFIXES[1]}"
        f"（收到 {suffix or '未知格式'}）"
    )


def _import_docx_text(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx，请先安装插件依赖") from exc

    document = docx.Document(str(path))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text.strip() for cell in row.cells)

    return "\n".join(line for line in lines if line)


def _import_xlsx_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("缺少 openpyxl，无法读取考纲") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row or ():
                    text = str(cell).strip() if cell is not None else ""
                    if text:
                        lines.append(text)
    finally:
        workbook.close()
    return "\n".join(lines)


def _import_docx_examples(path: Path) -> list[dict[str, Any]]:
    """从平台 Word 模板结构的文档里回读样例题（仅格式参考）。"""
    text = _import_docx_text(path)
    return parse_question_text_block(text)


def parse_question_text_block(text: str) -> list[dict[str, Any]]:
    """
    解析形如下面的题目文本块（Word 模板的正文格式）：

        CorelDraw窗口界面的组成（）(单选)【一般】【单选】
        A.标题栏、菜单栏
        B.属性栏
        答案：【B】
    """
    examples: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("知识点"):
            continue

        if line.startswith("答案"):
            if current is not None:
                answer = line.split("：")[-1] if "：" in line else line
                current["answer"] = "".join(
                    ch for ch in answer.upper() if ch in OPTION_LETTERS
                )
                examples.append(current)
                current = None
            continue

        if len(line) > 1 and line[0] in OPTION_LETTERS and line[1] in (".", "．"):
            if current is not None:
                current["options"].append(line[2:].strip())
            continue

        # 新题干
        if current is not None:
            examples.append(current)

        question_type = "单选"
        difficulty = "一般"
        if "【" in line:
            parts = [part for part in line.split("【") if part]
            tags = [part.split("】")[0] for part in parts if "】" in part]
            for tag in tags:
                if tag in ("单选", "多选", "判断"):
                    question_type = tag
                elif tag in ("简单", "一般", "困难"):
                    difficulty = tag

        current = {
            "stem": line,
            "type": question_type,
            "difficulty": difficulty,
            "options": [],
            "answer": "",
        }

    if current is not None:
        examples.append(current)

    return [item for item in examples if item.get("stem")]


def build_candidate_knowledge_prompt(syllabus_text: str, course_name: str = "") -> str:
    """
    构造［AI 提取候选］的提示词（FR2.4）。

    返回给 AI 的指令，要求输出候选知识点名称列表（8~20 个）。
    """
    trimmed = (syllabus_text or "").strip()
    if len(trimmed) > 6000:
        trimmed = trimmed[:6000]

    header = f"课程：{course_name}\n" if course_name else ""
    return (
        f"{header}以下是一份课程考纲/大纲的文本内容。\n"
        "请从中提取适合作为【课堂评估题】出题范围的知识点名称，"
        "每个名称应简短、粒度适中（不要整章、不要过细的知识碎片）。\n"
        "输出严格 JSON：{\"knowledge_points\": [\"名称1\", \"名称2\", ...]}，"
        "共 8~20 个，不要额外解释。\n\n"
        f"考纲内容：\n{trimmed}"
    )


def parse_candidate_knowledge_response(payload: Any) -> list[str]:
    """从 AI 响应里取出候选知识点名称列表。"""
    if isinstance(payload, str):
        from .ai_client import parse_json_content

        payload = parse_json_content(payload)
    if not isinstance(payload, dict):
        return []

    raw = payload.get("knowledge_points")
    if isinstance(raw, dict):
        raw = list(raw.keys())
    if not isinstance(raw, list):
        return []

    names: list[str] = []
    seen: set[str] = set()
    for item in raw:
        text = str(item).strip()
        if text and text not in seen:
            seen.add(text)
            names.append(text)
    return names


def question_from_example(example: dict[str, Any]) -> Question:
    """把样例题字典转成 :class:`Question`（仅用于展示/对照，不参与导出）。"""
    return Question(
        stem=str(example.get("stem") or ""),
        question_type=str(example.get("type") or "单选"),
        difficulty=str(example.get("difficulty") or "一般"),
        options=list(example.get("options") or []),
        answer=str(example.get("answer") or "").upper(),
    )
